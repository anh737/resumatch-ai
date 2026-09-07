"""Text extraction from uploaded CV/JD files (PDF, DOCX, TXT, MD).

The offline notebooks start from a CSV corpus / crawled HTML; the online path
starts from a single file an admin uploaded to MinIO. This module turns those
bytes into the plain text the structuring step (``core.structure``) expects.

PII scrubbing ports the regexes of ``pre-processing/preprocessing_cv.ipynb``
verbatim — apply :func:`scrub_pii` to resumes only, job postings keep their
contact/salary lines just like the offline JD pipeline.
"""

from __future__ import annotations

import io
import re
from pathlib import PurePosixPath

from utils.logger import get_logger

log = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# personal identifiers that survive the corpus anonymization (links, contacts, salary)
_PII_RES = [
    re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com/in/|github\.com/)\S+", re.I),
    re.compile(r"[\w.+-]+@[\w-]+\.\w+"),
    re.compile(r"\(?\d{3}\)?[-.\s]\d{3}[-.\s]\d{4}"),
    re.compile(r"(?:Starting|Ending)\s+Salary:\s*\$[\d,]+", re.I),
]

# unicode junk normalized away by the notebooks' clean()
_JUNK = [("\xa0", " "), ("​", ""), ("\xad", ""), ("－", "-")]


class UnsupportedFileError(ValueError):
    """The uploaded file's type cannot be turned into text."""


def clean_text(text: str) -> str:
    """Normalize unicode junk and collapse whitespace, keeping line structure."""
    for a, b in _JUNK:
        text = text.replace(a, b)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def scrub_pii(text: str) -> str:
    """Remove links, emails, phone numbers and salary lines (resumes only)."""
    for rx in _PII_RES:
        text = rx.sub(" ", text)
    return re.sub(r"[ \t]{2,}", " ", text).strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _plain_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def extract_text(data: bytes, filename: str) -> str:
    """Uploaded file bytes -> cleaned plain text.

    Raises :class:`UnsupportedFileError` for unknown extensions and
    ``ValueError`` when no text at all could be extracted (e.g. a scanned PDF
    with no text layer — OCR is out of scope for this service).
    """
    ext = PurePosixPath(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in {".txt", ".md"}:
        text = _plain_text(data)
    else:
        raise UnsupportedFileError(f"unsupported file type {ext or filename!r} (supported: {sorted(SUPPORTED_EXTENSIONS)})")

    text = clean_text(text)
    if not text:
        raise ValueError(f"no text could be extracted from {filename!r}")
    log.debug("extracted %d chars from %s", len(text), filename)
    return text
